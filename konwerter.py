
import json
from openpyxl import load_workbook
from flask import Flask, send_file, jsonify

INPUT_FILE = "produkty.xlsx"
OUTPUT_FILE = "data.json"

REQUIRED_COLUMNS = [
    "Kod",
    "Nazwa",
    "Opis",
    "Kod EAN",
    "Sztuki",
    "Jednostka Miary"
]

app = Flask(__name__)


def clean(value):
    if value is None:
        return ""

    return str(value).strip()


def generate_json():
    print("Otwieranie pliku Excel...")

    workbook = load_workbook(
        INPUT_FILE,
        read_only=True,
        data_only=True
    )

    sheet = workbook.active

    headers = {}

    for column, value in enumerate(sheet[1], start=1):
        header = clean(value.value)

        if header:
            headers[header] = column

    missing = [
        column
        for column in REQUIRED_COLUMNS
        if column not in headers
    ]

    if missing:
        print()
        print("BLAD: Brakuje kolumn:")

        for column in missing:
            print("-", column)

        print()
        print("Dostepne kolumny:")

        for column in headers:
            print("-", column)

        workbook.close()

        raise RuntimeError("Brakuje wymaganych kolumn w pliku Excel.")

    products = []

    for row in sheet.iter_rows(min_row=2, values_only=True):

        code = clean(row[headers["Kod"] - 1])

        if not code:
            continue

        name = clean(row[headers["Nazwa"] - 1])
        description = clean(row[headers["Opis"] - 1])
        ean = clean(row[headers["Kod EAN"] - 1])
        quantity = row[headers["Sztuki"] - 1]
        unit = clean(row[headers["Jednostka Miary"] - 1])

        if quantity is None or quantity == "":
            quantity = 0

        try:
            quantity = float(quantity)

            if quantity.is_integer():
                quantity = int(quantity)

        except (ValueError, TypeError):
            pass

        product = {
            "kod": code,
            "nazwa": name,
            "opis": description,
            "ean": ean,
            "sztuki": quantity,
            "jednostkaMiary": unit
        }

        products.append(product)

    result = {
        "products": products
    }

    with open(
        OUTPUT_FILE,
        "w",
        encoding="utf-8"
    ) as file:

        json.dump(
            result,
            file,
            ensure_ascii=False,
            indent=2
        )

    workbook.close()

    print()
    print("Gotowe!")
    print("Liczba produktow:", len(products))
    print("Utworzono plik:", OUTPUT_FILE)

    return len(products)


@app.get("/")
def home():
    return jsonify({
        "status": "ok",
        "message": "API magazynowe działa",
        "endpoint": "/data.json"
    })


@app.get("/data.json")
def data():
    return send_file(
        OUTPUT_FILE,
        mimetype="application/json"
    )


if __name__ == "__main__":
    # Najpierw wygeneruj aktualny JSON z Excela
    generate_json()

    # Następnie uruchom serwer HTTP
    app.run(
        host="0.0.0.0",
        port=8080
    )

